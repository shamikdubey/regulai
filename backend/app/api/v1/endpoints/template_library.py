"""
Template Library
  GET    /template-library/templates
  GET    /template-library/templates/{id}
  POST   /template-library/templates/generate
  POST   /template-library/templates/{id}/feedback
  PUT    /template-library/templates/{id}/verify
  GET    /template-library/templates/{id}/download
"""
import uuid
import json
import re
import io
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, field_validator
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import text

from app.db.database import get_db
from app.db.models import User
from app.services.auth_service import get_current_user
from app.core.config import get_settings

router = APIRouter()
settings = get_settings()

VALID_DOCUMENT_TYPES = (
    "CHECKLIST",
    "COVER_LETTER",
    "TECHNICAL_FILE",
    "REGULATORY_SUMMARY",
    "RISK_ASSESSMENT",
)

VALID_TRUST_LEVELS = ("VERIFIED", "HIGH", "MODERATE", "LOW", "UNVERIFIED")


# ── Pydantic models ────────────────────────────────────────────────────────────

class GenerateRequest(BaseModel):
    country: str
    domain: str
    document_type: str
    product_name: Optional[str] = None

    @field_validator("country", "domain")
    @classmethod
    def not_blank(cls, v):
        if not v or not v.strip():
            raise ValueError("field must not be blank")
        return v.strip()

    @field_validator("document_type")
    @classmethod
    def valid_doc_type(cls, v):
        if v not in VALID_DOCUMENT_TYPES:
            raise ValueError(f"document_type must be one of: {', '.join(VALID_DOCUMENT_TYPES)}")
        return v


class FeedbackRequest(BaseModel):
    rating: int
    is_accurate: Optional[bool] = None
    feedback_text: Optional[str] = None

    @field_validator("rating")
    @classmethod
    def rating_range(cls, v):
        if v < 1 or v > 5:
            raise ValueError("rating must be between 1 and 5")
        return v


class VerifyRequest(BaseModel):
    trust_score: int
    trust_level: str
    regulation_reference: Optional[str] = None

    @field_validator("trust_score")
    @classmethod
    def score_range(cls, v):
        if v < 0 or v > 100:
            raise ValueError("trust_score must be between 0 and 100")
        return v

    @field_validator("trust_level")
    @classmethod
    def valid_trust_level(cls, v):
        if v not in VALID_TRUST_LEVELS:
            raise ValueError(f"trust_level must be one of: {', '.join(VALID_TRUST_LEVELS)}")
        return v


# ── Helper: slugify ────────────────────────────────────────────────────────────

def _slugify(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r"[^\w\s-]", "", text)
    text = re.sub(r"[\s_-]+", "-", text)
    text = re.sub(r"^-+|-+$", "", text)
    return text[:120]


def _make_slug(country: str, domain: str, document_type: str, suffix: str = "") -> str:
    base = f"{country}-{domain}-{document_type}{'-' + suffix if suffix else ''}"
    return _slugify(base)


# ── Helper: html → plain text for DOCX ────────────────────────────────────────

def _strip_html(html: str) -> str:
    clean = re.sub(r"<[^>]+>", "", html or "")
    clean = re.sub(r"&nbsp;", " ", clean)
    clean = re.sub(r"&amp;", "&", clean)
    clean = re.sub(r"&lt;", "<", clean)
    clean = re.sub(r"&gt;", ">", clean)
    clean = re.sub(r"\n{3,}", "\n\n", clean)
    return clean.strip()


def _build_docx(template_row: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    title_para = doc.add_heading(template_row.get("title", "Template"), level=1)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata block
    meta = doc.add_paragraph()
    meta.add_run(f"Country: {template_row.get('country', '')}  |  "
                 f"Domain: {template_row.get('domain', '')}  |  "
                 f"Type: {template_row.get('document_type', '')}").italic = True

    if template_row.get("regulation_reference"):
        ref_para = doc.add_paragraph()
        ref_para.add_run("Regulatory Basis: ").bold = True
        ref_para.add_run(template_row["regulation_reference"])

    doc.add_paragraph()  # spacer

    # Content
    content = _strip_html(template_row.get("content_html") or "")
    for line in content.split("\n"):
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
        # Detect section headers (ALL CAPS or ending with colon)
        if line.isupper() or (line.endswith(":") and len(line) < 80):
            doc.add_heading(line.rstrip(":"), level=2)
        else:
            doc.add_paragraph(line)

    # Footer note
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run(
        f"Trust score: {template_row.get('trust_score', 0)}/100  |  "
        f"Trust level: {template_row.get('trust_level', 'UNVERIFIED')}  |  "
        f"Generated by RegulAI"
    ).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ── Endpoints ──────────────────────────────────────────────────────────────────

@router.get("/template-library/templates")
async def list_templates(
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
    country: Optional[str] = None,
    domain: Optional[str] = None,
    document_type: Optional[str] = None,
    trust_level: Optional[str] = None,
    search: Optional[str] = None,
):
    query = """
        SELECT id, tenant_id, title, slug, country, domain, document_type,
               regulation_reference, trust_score, trust_level,
               avg_user_rating, feedback_count, is_ai_generated,
               generation_model, created_at, updated_at
        FROM document_templates
        WHERE (tenant_id IS NULL OR tenant_id = :tenant_id)
    """
    params: dict = {"tenant_id": str(user.tenant_id)}

    if country:
        query += " AND country ILIKE :country"
        params["country"] = f"%{country}%"
    if domain:
        query += " AND domain = :domain"
        params["domain"] = domain
    if document_type:
        query += " AND document_type = :document_type"
        params["document_type"] = document_type
    if trust_level:
        query += " AND trust_level = :trust_level"
        params["trust_level"] = trust_level
    if search:
        query += " AND (title ILIKE :search OR regulation_reference ILIKE :search)"
        params["search"] = f"%{search}%"

    query += " ORDER BY trust_score DESC, avg_user_rating DESC"

    result = await db.execute(text(query), params)
    return [dict(r) for r in result.mappings().all()]


@router.get("/template-library/templates/{template_id}")
async def get_template(
    template_id: uuid.UUID,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(text("""
        SELECT * FROM document_templates
        WHERE id = :id AND (tenant_id IS NULL OR tenant_id = :tenant_id)
    """), {"id": str(template_id), "tenant_id": str(user.tenant_id)})
    row = result.mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Template not found")
    return dict(row)


@router.post("/template-library/templates/generate", status_code=201)
async def generate_template(
    body: GenerateRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    import anthropic as anthropic_sdk

    product_note = f"\nProduct (if applicable): {body.product_name}" if body.product_name else ""

    prompt = f"""You are a senior regulatory affairs expert creating a {body.document_type} template.

Country: {body.country}
Regulatory Domain: {body.domain}{product_note}

Create a detailed, professional {body.document_type} that:
1. References specific regulations by name and article number
2. Includes all mandatory sections required by {body.country} regulations
3. Uses [PLACEHOLDER] for fields the user must fill in
4. Is formatted clearly with numbered sections
5. Includes a "Regulatory Basis" section citing exact regulations

For the trust score, extract:
- regulation_reference: the primary regulation this is based on
- confidence_note: one sentence explaining basis for this template

IMPORTANT: Only reference regulations you are certain exist.
Mark uncertain items with [VERIFY: description]

Return JSON:
{{
  "title": "template title",
  "content_html": "full HTML content with sections using <h2>, <p>, <ol>, <ul> tags",
  "regulation_reference": "specific regulation cited",
  "confidence_note": "basis for this template",
  "suggested_trust_score": 40
}}"""

    client = anthropic_sdk.Anthropic(api_key=settings.ANTHROPIC_API_KEY)
    msg = client.messages.create(
        model=settings.CLAUDE_MODEL,
        max_tokens=3000,
        messages=[{"role": "user", "content": prompt}],
    )

    raw = msg.content[0].text.strip()
    if raw.startswith("```"):
        raw = "\n".join(raw.split("\n")[1:])
        if raw.endswith("```"):
            raw = raw[:-3]

    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        raise HTTPException(status_code=502, detail="AI returned invalid JSON")

    title = data.get("title") or f"{body.document_type} — {body.country} {body.domain}"
    content_html = data.get("content_html", "")
    regulation_reference = data.get("regulation_reference", "")
    suggested_score = int(data.get("suggested_trust_score", 40))
    trust_score = max(0, min(100, suggested_score))

    # Build unique slug
    base_slug = _make_slug(body.country, body.domain, body.document_type)
    suffix = str(uuid.uuid4())[:8]
    slug = f"{base_slug}-{suffix}"

    template_id = str(uuid.uuid4())

    await db.execute(text("""
        INSERT INTO document_templates (
            id, tenant_id, title, slug, country, domain, document_type,
            content_html, regulation_reference, trust_score, trust_level,
            is_ai_generated, generation_model
        ) VALUES (
            :id, :tenant_id, :title, :slug, :country, :domain, :document_type,
            :content_html, :regulation_reference, :trust_score, 'UNVERIFIED',
            TRUE, :model
        )
    """), {
        "id": template_id,
        "tenant_id": str(user.tenant_id),
        "title": title,
        "slug": slug,
        "country": body.country,
        "domain": body.domain,
        "document_type": body.document_type,
        "content_html": content_html,
        "regulation_reference": regulation_reference,
        "trust_score": trust_score,
        "model": settings.CLAUDE_MODEL,
    })
    await db.commit()

    return {
        "id": template_id,
        "title": title,
        "slug": slug,
        "country": body.country,
        "domain": body.domain,
        "document_type": body.document_type,
        "content_html": content_html,
        "regulation_reference": regulation_reference,
        "trust_score": trust_score,
        "trust_level": "UNVERIFIED",
        "is_ai_generated": True,
        "generation_model": settings.CLAUDE_MODEL,
        "confidence_note": data.get("confidence_note", ""),
    }


@router.post("/template-library/templates/{template_id}/feedback", status_code=201)
async def submit_feedback(
    template_id: uuid.UUID,
    body: FeedbackRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(text("""
        SELECT id FROM document_templates
        WHERE id = :id AND (tenant_id IS NULL OR tenant_id = :tenant_id)
    """), {"id": str(template_id), "tenant_id": str(user.tenant_id)})
    if not result.first():
        raise HTTPException(status_code=404, detail="Template not found")

    feedback_id = str(uuid.uuid4())
    await db.execute(text("""
        INSERT INTO template_feedback
            (id, template_id, user_id, tenant_id, rating, is_accurate, feedback_text)
        VALUES
            (:id, :template_id, :user_id, :tenant_id, :rating, :is_accurate, :feedback_text)
    """), {
        "id": feedback_id,
        "template_id": str(template_id),
        "user_id": str(user.id),
        "tenant_id": str(user.tenant_id),
        "rating": body.rating,
        "is_accurate": body.is_accurate,
        "feedback_text": body.feedback_text,
    })

    await db.execute(text("""
        UPDATE document_templates
        SET
            avg_user_rating = (
                SELECT ROUND(AVG(rating)::numeric, 2)
                FROM template_feedback
                WHERE template_id = :template_id
            ),
            feedback_count = feedback_count + 1,
            updated_at = now()
        WHERE id = :template_id
    """), {"template_id": str(template_id)})

    await db.commit()
    return {"id": feedback_id, "status": "recorded"}


@router.put("/template-library/templates/{template_id}/verify")
async def verify_template(
    template_id: uuid.UUID,
    body: VerifyRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    if getattr(user, "role", None) != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")

    result = await db.execute(text("""
        SELECT id FROM document_templates WHERE id = :id
    """), {"id": str(template_id)})
    if not result.first():
        raise HTTPException(status_code=404, detail="Template not found")

    await db.execute(text("""
        UPDATE document_templates
        SET
            trust_score = :trust_score,
            trust_level = :trust_level,
            regulation_reference = COALESCE(:regulation_reference, regulation_reference),
            verified_by = :verified_by,
            verified_at = now(),
            updated_at = now()
        WHERE id = :id
    """), {
        "trust_score": body.trust_score,
        "trust_level": body.trust_level,
        "regulation_reference": body.regulation_reference,
        "verified_by": str(user.id),
        "id": str(template_id),
    })
    await db.commit()
    return {"status": "verified", "trust_score": body.trust_score, "trust_level": body.trust_level}


@router.get("/template-library/templates/{template_id}/download")
async def download_template(
    template_id: uuid.UUID,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(text("""
        SELECT * FROM document_templates
        WHERE id = :id AND (tenant_id IS NULL OR tenant_id = :tenant_id)
    """), {"id": str(template_id), "tenant_id": str(user.tenant_id)})
    row = result.mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Template not found")

    docx_bytes = _build_docx(dict(row))
    safe_title = re.sub(r"[^\w\s-]", "", row.get("title", "template"))
    safe_title = re.sub(r"\s+", "_", safe_title)[:60]
    filename = f"{safe_title}.docx"

    return StreamingResponse(
        io.BytesIO(docx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
